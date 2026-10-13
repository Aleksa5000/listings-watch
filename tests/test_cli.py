"""The command line, driven with typer's CliRunner and fake sources (no network)."""

import csv
import io

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from helpers import FakeSource, items_for
from listings_watch import cli
from listings_watch.config import Settings

runner = CliRunner()


@pytest.fixture
def fake_sources(monkeypatch):
    shop = FakeSource("shop", items_for("shop", 4))
    monkeypatch.setattr(cli, "get_sources", lambda names=None: [shop])
    return shop


@pytest.fixture
def db(tmp_path, monkeypatch):
    monkeypatch.delenv("ALERT_WEBHOOK_URL", raising=False)
    return str(tmp_path / "cli.db")


def invoke(db, *args):
    return runner.invoke(cli.app, ["--db", db, *args])


def test_run_prints_a_summary_and_exits_zero(db, fake_sources):
    result = invoke(db, "run")
    assert result.exit_code == 0, result.output
    assert "Run #1" in result.output and "shop" in result.output and "new    4" in result.output
    assert "All sources healthy." in result.output


def test_run_exits_nonzero_when_a_source_is_broken(db, fake_sources):
    invoke(db, "run")
    fake_sources.items = []
    result = invoke(db, "run")
    assert result.exit_code == 1
    assert "BROKEN" in result.output and "0 items" in result.output
    assert "no alert" in result.output  # webhook not configured in this test


def test_run_with_unknown_source_is_a_usage_error(db):
    result = invoke(db, "run", "--source", "nope")
    assert result.exit_code == 2 and "unknown source" in result.output


def test_report_before_and_after_a_run(db, fake_sources):
    assert "No runs yet" in invoke(db, "report").output
    invoke(db, "run")
    out = invoke(db, "report").output
    assert "Latest run #1" in out and "found 4" in out and "shop" in out and "active" in out


def test_export_csv_and_xlsx(db, fake_sources, tmp_path):
    invoke(db, "run")
    target = tmp_path / "out" / "l.csv"
    result = invoke(db, "export", "--format", "csv", "--out", str(target))
    assert result.exit_code == 0 and "Exported 4 rows" in result.output
    assert len(list(csv.DictReader(io.StringIO(target.read_text(encoding="utf-8"))))) == 4
    xlsx = tmp_path / "l.xlsx"
    assert invoke(db, "export", "-f", "xlsx", "-o", str(xlsx)).exit_code == 0
    assert load_workbook(xlsx).active.max_row == 5


def test_export_filters_by_status(db, fake_sources, tmp_path):
    invoke(db, "run")
    fake_sources.items = fake_sources.items[:3]
    for _ in range(2):
        invoke(db, "run")
    target = tmp_path / "removed.csv"
    invoke(db, "export", "--status", "removed", "--out", str(target))
    rows = list(csv.DictReader(io.StringIO(target.read_text(encoding="utf-8"))))
    assert [r["external_id"] for r in rows] == ["id-4"]


def test_settings_from_env():
    s = Settings.from_env(
        {
            "LW_DELAY": "0.25",
            "LW_REMOVED_AFTER": "3",
            "ALERT_WEBHOOK_URL": " https://h.test/x ",
            "LW_MAX_PAGES": "",
        }
    )
    assert (s.delay, s.removed_after, s.max_pages) == (0.25, 3, 0)
    assert s.alert_webhook_url == "https://h.test/x"
    assert Settings.from_env({}).alert_webhook_url is None
    with pytest.raises(ValueError, match="LW_DELAY"):
        Settings.from_env({"LW_DELAY": "fast"})


def test_bad_env_value_is_reported_cleanly(db, monkeypatch):
    monkeypatch.setenv("LW_ATTEMPTS", "many")
    result = invoke(db, "report")
    assert result.exit_code == 2 and "LW_ATTEMPTS" in result.output

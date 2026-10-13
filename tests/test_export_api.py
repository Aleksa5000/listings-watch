"""Export formats and the HTTP API, on a small database built with the real pipeline."""

import asyncio
import csv
import io
from datetime import datetime

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from helpers import FakeSource, make
from listings_watch.api import create_app
from listings_watch.export import COLUMNS, to_csv, to_xlsx
from listings_watch.models import Listing
from listings_watch.runner import run_pipeline


def pipeline(settings, sessions, sources):
    return asyncio.run(run_pipeline(settings, sessions, sources, notify=lambda url, text: True))


@pytest.fixture
def populated(settings, sessions):
    shop_a = FakeSource(
        "a",
        [
            make("a", 1, title="Blue Widget", price=9.5, extra={"tags": ["x"], "n": 1}),
            make("a", 2, title="Red Gadget 100% off"),
            make("a", 3, title="Old Thing"),
        ],
    )
    shop_b = FakeSource("b", [make("b", 1, title="blue widget"), make("b", 2, title="Green Lamp")])
    pipeline(settings, sessions, [shop_a, shop_b])
    shop_a.items = shop_a.items[:2]
    for _ in range(2):  # Old Thing missed twice -> removed
        pipeline(settings, sessions, [shop_a, shop_b])
    return sessions


@pytest.fixture
def client(settings, engine, populated):
    return TestClient(create_app(settings, engine))


# --- export -------------------------------------------------------------------------------


def listing(**kw):
    return Listing(
        source="a",
        external_id=kw.pop("external_id", "1"),
        url="https://example.test/1",
        title=kw.pop("title", "T"),
        first_seen=datetime(2026, 1, 2, 3, 4, 5),
        **kw,
    )


def test_csv_has_header_json_extra_and_iso_dates():
    text = to_csv([listing(price=4.5, currency="GBP", extra={"b": 1, "a": [1, 2]}, title="x, y")])
    rows = list(csv.reader(io.StringIO(text)))
    assert rows[0] == COLUMNS
    row = dict(zip(COLUMNS, rows[1], strict=True))
    assert row["title"] == "x, y" and row["price"] == "4.5" and row["currency"] == "GBP"
    assert row["first_seen"] == "2026-01-02T03:04:05" and row["status"] == "active"
    assert row["extra"] == '{"a": [1, 2], "b": 1}' and row["duplicate_of"] == ""


@pytest.mark.parametrize("evil", ['=HYPERLINK("http://evil","x")', "+1+1", "-2+3", "@SUM(A1)"])
def test_scraped_text_cannot_become_a_spreadsheet_formula(evil):
    csv_cell = list(csv.reader(io.StringIO(to_csv([listing(title=evil)]))))[1][2]
    assert csv_cell == "'" + evil
    sheet = load_workbook(io.BytesIO(to_xlsx([listing(title=evil)]))).active
    cell = sheet.cell(row=2, column=3)
    assert cell.value == "'" + evil and cell.data_type == "s"  # stored as text, not "f"


def test_xlsx_opens_with_header_filter_and_values():
    data = to_xlsx([listing(price=4.5, title="Hello"), listing(external_id="2", title="World")])
    sheet = load_workbook(io.BytesIO(data)).active
    assert sheet.title == "listings"
    assert [c.value for c in sheet[1]] == COLUMNS and sheet[1][0].font.bold
    assert [sheet.cell(row=r, column=3).value for r in (2, 3)] == ["Hello", "World"]
    assert sheet.cell(row=2, column=4).value == 4.5
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref


def test_empty_export_is_still_a_valid_file():
    assert to_csv([]).strip() == ",".join(COLUMNS)
    assert load_workbook(io.BytesIO(to_xlsx([]))).active.max_row == 1


# --- API ----------------------------------------------------------------------------------


def test_listings_hides_duplicates_by_default(client):
    body = client.get("/listings").json()
    titles = sorted(i["title"] for i in body["items"])
    assert body["total"] == 4 and titles == [
        "Blue Widget",
        "Green Lamp",
        "Old Thing",
        "Red Gadget 100% off",
    ]


def test_listings_can_include_duplicates_and_show_where_they_point(client):
    body = client.get("/listings", params={"include_duplicates": True}).json()
    assert body["total"] == 5
    dup = next(i for i in body["items"] if i["title"] == "blue widget")
    original = next(i for i in body["items"] if i["title"] == "Blue Widget")
    assert dup["duplicate_of"] == original["id"] and dup["duplicate_score"] == 100.0


def test_listings_filters(client):
    assert (
        client.get("/listings", params={"status": "removed"}).json()["items"][0]["title"]
        == "Old Thing"
    )
    assert client.get("/listings", params={"status": "active"}).json()["total"] == 3
    assert {
        i["source"] for i in client.get("/listings", params={"source": "b"}).json()["items"]
    } == {"b"}
    assert client.get("/listings", params={"q": "GADGET"}).json()["total"] == 1
    assert client.get("/listings", params={"q": "100%"}).json()["total"] == 1
    assert client.get("/listings", params={"q": "nothing like this"}).json() == {
        "total": 0,
        "limit": 100,
        "offset": 0,
        "items": [],
    }


def test_listings_pagination_and_validation(client):
    page = client.get("/listings", params={"limit": 2, "offset": 3}).json()
    assert page["total"] == 4 and len(page["items"]) == 1 and page["limit"] == 2
    assert client.get("/listings", params={"status": "deleted"}).status_code == 422
    assert client.get("/listings", params={"limit": 0}).status_code == 422
    assert client.get("/listings", params={"limit": 5000}).status_code == 422


def test_listing_json_shape(client):
    item = client.get("/listings", params={"q": "blue widget", "source": "a"}).json()["items"][0]
    assert set(item) == {
        "id", "source", "external_id", "title", "price", "currency", "url", "extra",
        "first_seen", "last_seen", "status", "duplicate_of", "duplicate_score",
    }  # fmt: skip
    assert item["price"] == 9.5 and item["extra"] == {"tags": ["x"], "n": 1}


def test_csv_endpoint_respects_filters(client):
    response = client.get("/listings.csv", params={"status": "active"})
    assert response.headers["content-type"].startswith("text/csv")
    assert "listings.csv" in response.headers["content-disposition"]
    rows = list(csv.DictReader(io.StringIO(response.text)))
    assert sorted(r["title"] for r in rows) == ["Blue Widget", "Green Lamp", "Red Gadget 100% off"]


def test_xlsx_endpoint(client):
    response = client.get("/listings.xlsx")
    assert response.status_code == 200
    assert load_workbook(io.BytesIO(response.content)).active.max_row == 5  # header + 4


def test_runs_endpoint_lists_per_source_counts_newest_first(client):
    runs = client.get("/runs").json()
    assert [r["id"] for r in runs] == [3, 2, 1]
    first = {s["source"]: s for s in runs[-1]["sources"]}
    assert (first["a"]["new"], first["b"]["new"], first["b"]["duplicates"]) == (3, 2, 1)
    last = {s["source"]: s for s in runs[0]["sources"]}
    assert last["a"]["removed"] == 1 and last["a"]["status"] == "ok"
    assert client.get("/runs", params={"limit": 1}).json()[0]["id"] == 3


def test_health_ok_then_degraded_after_a_broken_source(client, settings, sessions):
    health = client.get("/health").json()
    assert health["status"] == "ok" and health["listings"] == {"active": 4, "removed": 1}
    assert health["last_run"]["id"] == 3
    pipeline(settings, sessions, [FakeSource("a", [])])
    assert client.get("/health").json()["status"] == "degraded"


def test_health_on_an_empty_database(settings, engine):
    body = TestClient(create_app(settings, engine)).get("/health").json()
    assert body["status"] == "no_runs" and body["last_run"] is None
    assert body["listings"] == {"active": 0, "removed": 0}
